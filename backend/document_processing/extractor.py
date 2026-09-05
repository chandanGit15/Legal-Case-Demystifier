"""Document text extraction. Images are not OCR'd locally — they are sent
to the AI service for transcription/analysis (which is why they get the
'ocr_pending' status until analyzed)."""
import os

from pypdf import PdfReader
from docx import Document as DocxDocument


def extract_text(file_type: str, path: str) -> tuple[str, str]:
    """Return (text, tool_used). tool_used is '' for images (OCR handled by AI)."""
    if file_type == "pdf":
        reader = PdfReader(path)
        pages = []
        for page in reader.pages:
            try:
                pages.append(page.extract_text() or "")
            except Exception:
                pages.append("")
        return "\n\n".join(pages), "pypdf"

    if file_type == "docx":
        doc = DocxDocument(path)
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts), "python-docx"

    if file_type == "txt":
        with open(path, "r", encoding="utf-8", errors="replace") as fh:
            return fh.read(), "plain text"

    # Images: OCR is performed by the AI service when an API key is configured.
    return "", ""


def file_type_of(filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower().lstrip(".")
    if ext == "pdf":
        return "pdf"
    if ext == "docx":
        return "docx"
    if ext == "txt":
        return "txt"
    return "image"